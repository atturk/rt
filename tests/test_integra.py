"""Uso Integra: contenuto della scaletta e della riscrittura, mai solo contesto."""
import io
import json
from unittest.mock import Mock

import fitz
import pytest

from rt.core import archive
from rt.core.lesson_paths import lesson_path
from rt.core.materials import material_context_parts, material_pages
from rt.core.models import Outline, OutlineMacro, OutlineUnit
from rt.core.segments import load_segments_json
from rt.core.source_selection import record_used_selection, selection, set_primary
from rt.core.sources import add_source, load_sources, remove_source, update_source
from rt.llm.client import LLMClient
from rt.llm.prompts import INTEGRA_OUTLINE_RULE, REWRITE_SYSTEM_PROMPT
from rt.pipeline.material_mapping import assign_material_pages
from rt.pipeline.outline import load_outline
from rt.pipeline.source_update import (
    _read_state, prepare_source_update, run_source_update, segment_source_ids, sources_pending,
)
from tests.test_source_update import legacy


def paged_pdf(label='Argomento'):
    with fitz.open() as document:
        for number in range(1, 6):
            page = document.new_page()
            page.insert_text((50, 50), f'{label} {number}: contenuto specifico della pagina {number}.')
        return document.tobytes()


def selected_integra(lesson):
    source = add_source(lesson, 'materiale', 'integra.pdf', paged_pdf(),
                        pagine=[2, 3, 4, 5], uso='integra', estrai='immagini')
    archive.save_meta(source.impronta, {'ignorate': [3], 'eliminate': [4]})
    return source


def one_unit_per_segment(segments):
    return Outline(lesson_title='Lezione', macro_sections=[OutlineMacro(
        id='1', title='Contenuti', units=[OutlineUnit(
            id=f'1.{i}', title=f'Parte {i}', start_segment_id=s.id, end_segment_id=s.id)
            for i, s in enumerate(segments, 1)])])


def test_pdf_integra_selected_pages_follow_audio_and_keep_timestamps(tmp_path):
    legacy(tmp_path)
    lesson = str(tmp_path)
    source = selected_integra(lesson)
    notes = add_source(lesson, 'appunti', 'ordinary.txt', b'Appunti ordinari')
    result = prepare_source_update(lesson, mock=True)['segments']
    assert [s.source_file for s in result] == ['src_000001', notes.id, source.id, source.id]
    assert [s.text_raw for s in result[-2:]] == [p['testo'] for p in material_pages(lesson, source)]
    assert [p['numero'] for p in material_pages(lesson, source)] == [2, 5]
    assert all(s.integra and s.fonte == 'integra.pdf' and s.flags == ['materiale'] for s in result[-2:])
    assert result[0].start_seconds == 0 and result[0].end_seconds == 1
    assert result[-1].start_seconds == 2 and result[-1].end_seconds == 2.001
    saved = load_segments_json(lesson_path(lesson, 'segments.json'))
    assert saved.segments == result and saved.audio_duration_seconds == 2
    assert prepare_source_update(lesson, mock=True)['segments'] == result
    assert not sources_pending(lesson)


def test_classifier_only_handles_context_and_integra_images_follow_own_segments(tmp_path, monkeypatch):
    legacy(tmp_path)
    lesson = str(tmp_path)
    source = selected_integra(lesson)
    context = add_source(lesson, 'materiale', 'context.txt', b'Contesto ordinario', uso='contesto')
    result = prepare_source_update(lesson, mock=True)['segments']
    classifier = Mock(return_value=(['1.1'], []))
    monkeypatch.setattr('rt.pipeline.material_mapping.place_page_in_outline', classifier)
    state = assign_material_pages(lesson, one_unit_per_segment(result), mock=True)
    assert classifier.call_count == 1
    assert classifier.call_args.args[3].id == context.id
    assert {item['source_id'] for items in state['pages_by_unit'].values() for item in items} == {context.id}
    from rt.pipeline.add_images import load_image_descriptions
    descriptions = load_image_descriptions(lesson)
    for uid, number in [('1.2', 2), ('1.3', 5)]:
        hashes = state['units'][uid]
        assert hashes and all(descriptions[h]['source'] == f'material:{source.id}:{number}' for h in hashes)
    assert '1.1' not in state['units']
    parts = material_context_parts(lesson, mock=True)
    assert len(parts) == 1 and 'Contesto ordinario' in parts[0]
    assert 'integra.pdf' not in parts[0]


def test_material_primary_and_integra_images_do_not_mix_sources(tmp_path):
    (tmp_path / 'info.yaml').write_text('materia: Test\n')
    lesson = str(tmp_path)
    primary = add_source(lesson, 'materiale', 'primary.pdf', paged_pdf('Principale'), pagine=[1, 3], estrai='immagini')
    set_primary(lesson, primary.id)
    integra = selected_integra(lesson)
    segments = prepare_source_update(lesson, mock=True)['segments']
    state = assign_material_pages(lesson, one_unit_per_segment(segments), mock=True)
    from rt.pipeline.add_images import load_image_descriptions
    descriptions = load_image_descriptions(lesson)
    for index, (source, number) in enumerate([(primary, 1), (primary, 3), (integra, 2), (integra, 5)], 1):
        hashes = state['units'][f'1.{index}']
        assert hashes and all(descriptions[h]['source'] == f'material:{source.id}:{number}' for h in hashes)


@pytest.mark.parametrize('kind', ['appunti', 'audio', 'txt', 'docx'])
def test_all_integra_types_join_explicit_primary(tmp_path, kind):
    (tmp_path / 'info.yaml').write_text('materia: Test\n')
    lesson = str(tmp_path)
    primary = add_source(lesson, 'appunti', 'primary.txt', b'Fonte principale')
    set_primary(lesson, primary.id)
    if kind == 'audio':
        from tests.test_source_update import write_audio
        write_audio(tmp_path / 'incoming.wav')
        content = (tmp_path / 'incoming.wav').read_bytes()
        filename, tipo = 'incoming.wav', 'audio'
    elif kind == 'docx':
        from docx import Document
        document = Document()
        document.add_heading('Nuovo argomento', 1)
        document.add_paragraph('Da coprire integralmente')
        stream = io.BytesIO()
        document.save(stream)
        content, filename, tipo = stream.getvalue(), 'new.docx', 'materiale'
    else:
        content, filename = b'Nuovo argomento\n\nDa coprire integralmente', 'new.txt'
        tipo = 'appunti' if kind == 'appunti' else 'materiale'
    source = add_source(lesson, tipo, filename, content, uso='integra')
    segments = prepare_source_update(lesson, mock=True)['segments']
    assert segments[0].source_file == primary.id and not segments[0].integra
    assert len(segments) > 1
    assert all(s.source_file == source.id and s.integra and s.fonte for s in segments[1:])
    assert material_context_parts(lesson, mock=True) == []
    remove_source(lesson, source.id)
    assert source.id not in segment_source_ids(load_sources(lesson, include_removed=True), _read_state(lesson), primary.id)


def test_integra_prompts_and_extend_keep_existing_outline_and_draft(tmp_path, monkeypatch):
    legacy(tmp_path)
    lesson = str(tmp_path)
    monkeypatch.setattr('rt.services.classifier.classify_units', lambda *a, **k: None)
    run_source_update(lesson, mode='rebuild', mock=True)
    old_outline = load_outline(lesson)
    with open(lesson_path(lesson, 'draft.json')) as stream:
        old_draft = json.load(stream)
    source = selected_integra(lesson)
    calls = []
    original = LLMClient.call_structured

    def call(self, **kwargs):
        calls.append(kwargs)
        return original(self, **kwargs)

    monkeypatch.setattr(LLMClient, 'call_structured', call)
    result = run_source_update(lesson, mode='extend', mock=True)
    assert result['status'] == 'draft_validated'
    extended = load_outline(lesson)
    old_units = [u for m in old_outline.macro_sections for u in m.units]
    new_units = [u for m in extended.macro_sections for u in m.units]
    assert new_units[:len(old_units)] == old_units
    assert len(new_units) > len(old_units)
    with open(lesson_path(lesson, 'draft.json')) as stream:
        draft = json.load(stream)
    assert draft['units'][:len(old_draft['units'])] == old_draft['units']
    rewrite_calls = [c for c in calls if c['job_name'] == 'rewrite']
    assert rewrite_calls and all('[integra: integra.pdf]' in c['prompt'] for c in rewrite_calls)
    assert all('[integra: …]' in c['system_prompt'] for c in rewrite_calls)
    assert not set(result['changed_units']) & {u.id for u in old_units}
    assert _read_state(lesson)['usi_usati'][source.id] == 'integra'

    # Ricostruire usa il sommario marcato e la regola anche nel prompt della scaletta.
    calls.clear()
    run_source_update(lesson, mode='rebuild', mock=True)
    outline_calls = [c for c in calls if c['job_name'] == 'outline' and '[seg_' in c['prompt']]
    assert outline_calls and '[integra: integra.pdf]' in outline_calls[0]['prompt']
    assert INTEGRA_OUTLINE_RULE in outline_calls[0]['system_prompt']
    assert 'Questa regola non vale per i segmenti marcati [integra: …]' in REWRITE_SYSTEM_PROMPT


@pytest.mark.parametrize(('before', 'after'), [('contesto', 'integra'), ('integra', 'contesto'), (None, 'integra')])
def test_changing_used_integra_requires_rebuild_until_success(tmp_path, monkeypatch, before, after):
    legacy(tmp_path)
    lesson = str(tmp_path)
    monkeypatch.setattr('rt.services.classifier.classify_units', lambda *a, **k: None)
    source = add_source(lesson, 'materiale', 'new.txt', b'Nuovo tema', uso=before)
    run_source_update(lesson, mode='rebuild', mock=True)
    assert not selection(lesson)['solo_ricostruzione']
    update_source(lesson, source.id, uso=after)
    assert sources_pending(lesson) and selection(lesson)['solo_ricostruzione']
    with pytest.raises(ValueError, match='Ricostruisci'):
        run_source_update(lesson, mode='extend', mock=True)
    update_source(lesson, source.id, uso=before)
    assert not selection(lesson)['solo_ricostruzione']
    update_source(lesson, source.id, uso=after)
    run_source_update(lesson, mode='rebuild', mock=True)
    assert not sources_pending(lesson) and not selection(lesson)['solo_ricostruzione']
    assert _read_state(lesson)['usi_usati'][source.id] == after


def test_unused_integra_can_change_use_without_rebuild(tmp_path):
    legacy(tmp_path)
    lesson = str(tmp_path)
    record_used_selection(lesson)
    source = add_source(lesson, 'materiale', 'new.txt', b'Nuovo tema', uso='contesto')
    update_source(lesson, source.id, uso='integra')
    assert sources_pending(lesson) and not selection(lesson)['solo_ricostruzione']
