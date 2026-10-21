import io
from pathlib import Path
from unittest.mock import Mock

import pytest

from rt.core.materials import material_pages, page_count
from rt.core.source_options import parse_pages
from rt.core.sources import add_source
from rt.core.archive import load_meta
from tests.test_archive import lesson
from tests.test_materials import pdf_bytes


@pytest.mark.parametrize(('value', 'expected'), [('31-64', list(range(31, 65))), ('1,10,12-14', [1, 10, 12, 13, 14])])
def test_parse_pages(value, expected):
    assert parse_pages(value, 64) == expected


@pytest.mark.parametrize('value', ['0', '2-1', '1-65', '-1', 'foo', '1,', [0], [True]])
def test_invalid_pages(value):
    with pytest.raises(ValueError):
        parse_pages(value, 64)


def test_text_pdf_and_cp1252_without_vision(tmp_path, monkeypatch):
    directory = lesson(tmp_path, 'text')
    monkeypatch.setattr('rt.core.materials._describe', Mock(side_effect=AssertionError('vision')))
    source = add_source(directory, 'materiale', 'file.pdf', pdf_bytes(), estrai='testo')
    pages = material_pages(directory, source)
    assert pages[0]['testo'] and not pages[0]['immagini']
    text = add_source(directory, 'materiale', 'file.txt', 'Una città e un caffè'.encode('cp1252'))
    assert text.estrai == 'testo'
    assert material_pages(directory, text)[0]['testo'] == 'Una città e un caffè'


def test_docx_heading_table_and_pages(tmp_path, monkeypatch):
    from docx import Document
    document = Document()
    document.add_heading('Introduzione', 1)
    document.add_paragraph('Paragrafo prima della tabella')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Colonna uno'
    table.cell(0, 1).text = 'Colonna due'
    document.add_heading('Conclusione', 1)
    document.add_paragraph('Ultimo paragrafo')
    stream = io.BytesIO()
    document.save(stream)
    directory = lesson(tmp_path, 'word')
    source = add_source(directory, 'materiale', 'file.docx', stream.getvalue())
    pages = material_pages(directory, source)
    assert len(pages) == 2
    assert pages[0]['testo'].index('Paragrafo') < pages[0]['testo'].index('Colonna uno')
    assert 'Conclusione' in pages[1]['testo']
    assert not any(p['immagini'] for p in pages)
    assert page_count(str(Path(directory) / source.file)) == 2


def test_creation_options(tmp_path, monkeypatch):
    from rt.services.materials_lesson_service import create_source_lesson
    from rt.core.source_selection import selection
    from rt.core.sources import load_sources
    material = tmp_path / 'file.txt'
    material.write_text('Contenuto da studiare')
    directory = create_source_lesson(audio=[], materiali=[str(material)], date='2026-10-12',
        materia='Test', dest_dir=str(tmp_path), opzioni={'file.txt': {'pagine': '1', 'estrai': 'testo', 'uso': 'guida'}})
    source = load_sources(directory)[0]
    assert source.pagine == [1] and source.estrai == 'testo'
    assert selection(directory)['principale'] == source.id


def test_source_options_api_and_lookup(api_client, rt_db, tmp_path, monkeypatch):
    from tests.api_support import isolated_workspace, make_lesson
    from rt.core import archive
    import hashlib
    root = isolated_workspace(tmp_path, monkeypatch)
    make_lesson(root)
    lid = api_client.get('/api/v1/lessons').json()[0]['id']
    url = f'/api/v1/lessons/{lid}/sources'
    invalid = api_client.post(url, data={'tipo': 'materiale', 'pagine': '2'}, files={'file': ('file.pdf', pdf_bytes())})
    assert invalid.status_code == 422
    result = api_client.post(url, data={'tipo': 'materiale', 'uso': 'guida', 'estrai': 'testo'}, files={'file': ('file.txt', b'Contenuto')})
    assert result.status_code == 201, result.text
    source = result.json()
    assert api_client.get(url).json()['principale'] == source['id']
    changed = api_client.patch(f"{url}/{source['id']}", json={'pagine': '1', 'uso': 'integra'})
    assert changed.status_code == 200 and changed.json()['pagine'] == [1]
    assert api_client.get(url).json()['has_unused_sources']
    assert api_client.patch(f"{url}/{source['id']}", json={'pagine': '9'}).status_code == 422
    alias = hashlib.sha256(b'alias').hexdigest()
    archive.register_file(archive.original_path(source['impronta']), nome='file.txt', tipo='testo', alias=[alias])
    looked_up = api_client.post('/api/v1/archive/lookup', json={'sha': [alias, 'f' * 64]})
    assert looked_up.status_code == 200, looked_up.text
    assert looked_up.json()[alias]['impronta'] == source['impronta']
    assert looked_up.json()['f' * 64] is None


def test_text_mode_hides_images_in_shared_full_cache(tmp_path):
    directory = lesson(tmp_path, 'shared')
    source = add_source(directory, 'materiale', 'document.pdf', pdf_bytes())
    assert material_pages(directory, source)[0]['immagini']
    source = source.model_copy(update={'estrai': 'testo'})
    assert material_pages(directory, source)[0]['immagini'] == []
    from rt.core.archive import pages_for_use
    assert pages_for_use(source.impronta)[0]['immagini']


def test_docx_embedded_images_in_paragraphs_and_tables(tmp_path):
    from docx import Document
    from tests.test_materials import png_bytes
    document = Document()
    document.add_paragraph('Testo completo di una pagina con illustrazioni incorporate e tabella.')
    document.add_picture(io.BytesIO(png_bytes()))
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(png_bytes()))
    stream = io.BytesIO()
    document.save(stream)
    directory = lesson(tmp_path, 'word-images')
    source = add_source(directory, 'materiale', 'images.docx', stream.getvalue(), estrai='immagini')
    pages = material_pages(directory, source, mock=True)
    assert len(pages[0]['immagini']) == 2
