"""Testo e pagine immagine derivati una volta sola da una fonte immutabile."""
import base64
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
from contextlib import contextmanager

from rt.core.idempotency import compute_file_sha256
from rt.core.lesson_lock import lesson_lock
from rt.storage import fs

MATERIAL_EXTENSIONS = {'.pdf', '.png', '.jpg', '.jpeg', '.heic', '.pptx'}


def page_count(path):
    with fs.materialize(path) as real:
        suffix = os.path.splitext(real)[1].lower()
        if suffix == '.pdf':
            import fitz
            with fitz.open(real) as document:
                return len(document)
        if suffix == '.pptx':
            from pptx import Presentation
            return len(Presentation(real).slides)
        from PIL import Image
        from pillow_heif import register_heif_opener
        register_heif_opener()
        with Image.open(real) as image:
            image.verify()
        return 1


def _png_image(data):
    from PIL import Image, ImageOps
    try:
        from pillow_heif import register_heif_opener
        register_heif_opener()
    except ImportError:
        pass
    with Image.open(io.BytesIO(data)) as image:
        out = io.BytesIO()
        ImageOps.exif_transpose(image).convert('RGB').save(out, format='PNG')
        return out.getvalue()


def _describe(data, lesson_dir, mock):
    from rt.llm.client import LLMClient
    from rt.llm.prompts import ImageDescription, IMAGE_DESCRIPTION_SYSTEM_PROMPT_NO_CONTEXT
    from rt.services.prompt_settings import effective_system
    result = LLMClient(force_mock=mock).call_structured(
        prompt='Trascrivi tutto il testo leggibile e descrivi i contenuti della pagina per studiarli.',
        system_prompt=effective_system('image_description', IMAGE_DESCRIPTION_SYSTEM_PROMPT_NO_CONTEXT),
        response_model=ImageDescription, job_name='image_description', lesson_dir=lesson_dir,
        image_data_url='data:image/png;base64,' + base64.b64encode(data).decode('ascii'))
    return '\n'.join(str(value) for value in (result.slide_title, result.ocr_text, result.visual_elements) if value)


def _soffice():
    """LibreOffice, se installato: nel PATH o nella cartella standard di Windows."""
    found = shutil.which('soffice') or shutil.which('libreoffice')
    if found or os.name != 'nt':
        return found
    for base in (os.environ.get('ProgramFiles'), os.environ.get('ProgramFiles(x86)')):
        candidate = os.path.join(base or '', 'LibreOffice', 'program', 'soffice.exe')
        if base and os.path.isfile(candidate):
            return candidate
    return None


def _slide_command(real, target):
    """Quick Look su macOS, altrimenti LibreOffice; None se nessuno dei due c'è."""
    quicklook = shutil.which('qlmanage') if sys.platform == 'darwin' else None
    if quicklook:
        return [quicklook, '-p', '-o', target, real]
    soffice = _soffice()
    if soffice:
        return [soffice, '--headless', '--convert-to', 'pdf', '--outdir', target, real]
    return None


@contextmanager
def _quicklook_slides(real, count):
    """La resa delle slide (Quick Look o LibreOffice) è facoltativa; si usa solo se copre tutte
    le slide. Restituisce una funzione che rende una slide alla volta, così non tiene in
    memoria tutte le immagini."""
    with tempfile.TemporaryDirectory(prefix='rt-slides-') as target:
        command = _slide_command(real, target)
        if not command:
            yield None
            return
        render = None
        document = None
        try:
            subprocess.run(command, capture_output=True, check=True, timeout=120)
            from pathlib import Path
            import fitz
            for path in Path(target).rglob('*.pdf'):
                candidate = fitz.open(path)
                if len(candidate) == count:
                    document = candidate
                    render = lambda i: document[i].get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes('png')
                    break
                candidate.close()
            if render is None:
                images = sorted(Path(target).rglob('*.png'))
                if len(images) == count:
                    render = lambda i: images[i].read_bytes()
        except (OSError, subprocess.SubprocessError, ValueError, RuntimeError):
            render = None
        try:
            yield render
        finally:
            if document is not None:
                document.close()


def _slide_content(slide):
    texts, images = [], []
    def visit(shapes):
        for shape in shapes:
            if getattr(shape, 'has_text_frame', False):
                texts.append(shape.text)
            if getattr(shape, 'has_table', False):
                texts.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
            if hasattr(shape, 'image'):
                images.append(shape.image.blob)
            if hasattr(shape, 'shapes'):
                visit(shape.shapes)
    visit(slide.shapes)
    return '\n'.join(texts).strip(), images


def _extract(real):
    suffix = os.path.splitext(real)[1].lower()
    if suffix == '.pdf':
        import fitz
        from rt.core.material_page_filter import pdf_pages_have_figures
        with fitz.open(real) as document:
            figures = pdf_pages_have_figures(document)
            for page, has_figures in zip(document, figures):
                yield page.get_text().strip(), [page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes('png')], has_figures
    elif suffix == '.pptx':
        from pptx import Presentation
        presentation = Presentation(real)
        with _quicklook_slides(real, len(presentation.slides)) as render:
            for i, slide in enumerate(presentation.slides):
                text, pictures = _slide_content(slide)
                yield text, [render(i)] if render else [_png_image(data) for data in pictures], bool(pictures)
    else:
        with open(real, 'rb') as stream:
            yield '', [_png_image(stream.read())], True


def cached_pages(lesson_dir, source):
    path = os.path.join(lesson_dir, 'materials', source.impronta, 'pages.json')
    if not fs.isfile(path):
        return None
    with fs.open(path, encoding='utf-8') as stream:
        data = json.load(stream)
    return data['pages'] if data.get('impronta') == source.impronta else None


def material_pages(lesson_dir, source, *, mock=False, ctx=None):
    """Estrazione e Vision senza lock della lezione (possono durare minuti); il lock copre
    solo le scritture. La cache è per pagina e per impronta, quindi due lavori concorrenti
    sulla stessa fonte scrivono lo stesso contenuto."""
    path = os.path.join(lesson_dir, source.file)
    if compute_file_sha256(path) != source.impronta:
        raise ValueError('Una fonte originale è cambiata: aggiungila come nuova fonte')
    cached = cached_pages(lesson_dir, source)
    if cached is not None and all('has_figures' in page for page in cached):
        return cached
    folder = os.path.join(lesson_dir, 'materials', source.impronta)
    with lesson_lock(lesson_dir):
        fs.makedirs(folder, exist_ok=True)
    pages = []
    # Cache per pagina: una ripresa dopo errore non ripete le chiamate Vision riuscite.
    with fs.materialize(path) as real:
        from rt.core.material_page_filter import cleaned_page_texts, few_words
        extracted = list(_extract(real))
        cleaned = cleaned_page_texts([text for text, _, _ in extracted])
        for number, (text, images, has_figures) in enumerate(extracted, 1):
            if ctx:
                ctx.check_cancelled()
            entry_path = os.path.join(folder, f'page-{number}.json')
            if fs.isfile(entry_path):
                with fs.open(entry_path, encoding='utf-8') as stream:
                    saved = json.load(stream)
                if 'has_figures' in saved:
                    pages.append(saved)
                    continue
            mapping_text = cleaned[number - 1]
            sparse = few_words(mapping_text)
            if sparse and has_figures:
                described = '\n'.join(_describe(data, lesson_dir, mock) for data in images)
                text = '\n'.join(filter(None, [text, described]))
                mapping_text = described
            filenames = []
            with lesson_lock(lesson_dir):
                for index, data in enumerate(images, 1):
                    relative = os.path.join('materials', source.impronta, f'page-{number}-{index}.png')
                    with fs.open(os.path.join(lesson_dir, relative), 'wb') as stream:
                        stream.write(data)
                    filenames.append(relative)
                entry = {'numero': number, 'testo': text, 'immagini': filenames,
                         'mapping_text': mapping_text, 'has_figures': has_figures,
                         'mapping_skip': sparse and not has_figures}
                with fs.open(entry_path + '.tmp', 'w', encoding='utf-8') as stream:
                    json.dump(entry, stream, ensure_ascii=False)
                fs.replace(entry_path + '.tmp', entry_path)
            pages.append(entry)
    if not pages:
        raise ValueError('Il materiale non ha pagine')
    cache_path = os.path.join(folder, 'pages.json')
    with lesson_lock(lesson_dir):
        with fs.open(cache_path + '.tmp', 'w', encoding='utf-8') as stream:
            json.dump({'impronta': source.impronta, 'pages': pages}, stream, ensure_ascii=False)
        fs.replace(cache_path + '.tmp', cache_path)
    return pages


def material_context_parts(lesson_dir, *, mock=False, ctx=None, include_materials: bool = True):
    """Contesto per la riscrittura: solo le fonti che non sono già nei segmenti, una volta sola."""
    from rt.core.sources import load_sources
    from rt.core.source_selection import selection
    from rt.pipeline.source_update import _read_state, segment_source_ids
    sources = load_sources(lesson_dir)
    parts = []
    primary = selection(lesson_dir)['principale']
    state = _read_state(lesson_dir)
    cache = state.get('sources', {})
    in_segments = segment_source_ids(sources, state, primary)
    for source in sources:
        if source.id in in_segments:
            continue
        if source.tipo == 'materiale':
            if include_materials:
                for page in material_pages(lesson_dir, source, mock=mock, ctx=ctx):
                    parts.append(f"[{source.id}, {os.path.basename(source.file)}, pagina {page['numero']}]\n{page['testo']}")
        else:
            texts = [segment['text_raw'] for segment in cache.get(source.id, {}).get('segments', [])]
            if texts:
                parts.append(f'[{source.id}, {os.path.basename(source.file)}]\n' + '\n'.join(texts))
    return parts
